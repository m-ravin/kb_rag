"""
Extraction stage — pulls raw text out of uploaded PDF/DOCX/PPTX files.

Pure byte-content I/O with no Azure SDK clients involved, so this module is
fully unit-testable without any environment configuration or network access.
"""

import os
import tempfile
from typing import Any

import fitz  # PyMuPDF — reads PDFs
from docx import Document as DocxDocument
from pptx import Presentation

_MAGIC_BYTES: dict[str, bytes] = {
    "pdf": b"%PDF",
    "docx": b"PK\x03\x04",  # ZIP-based Office format
    "pptx": b"PK\x03\x04",
}


def validate_magic_bytes(raw_bytes: bytes, filename: str) -> None:
    """Reject files whose magic bytes don't match the declared extension."""
    ext = filename.rsplit(".", 1)[-1].lower()
    expected = _MAGIC_BYTES.get(ext)
    if expected and not raw_bytes.startswith(expected):
        raise ValueError(
            f"File content does not match extension .{ext} — upload rejected."
        )


def extract_content(raw_bytes: bytes, filename: str) -> tuple[str, dict]:
    """
    Reads a document file and pulls out all the text.
    Like a person reading a book and typing out every word.
    """
    ext = filename.rsplit(".", 1)[-1].lower()
    metadata: dict[str, Any] = {"filename": filename, "file_type": ext}

    if ext == "pdf":
        return _extract_pdf(raw_bytes, metadata)
    elif ext == "docx":
        return _extract_docx(raw_bytes, metadata)
    elif ext in ("pptx", "ppt"):
        return _extract_pptx(raw_bytes, metadata)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(raw_bytes: bytes, metadata: dict) -> tuple[str, dict]:
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(raw_bytes)
        tmp_path = tmp.name
    try:
        doc = fitz.open(tmp_path)
        pages = [page.get_text() for page in doc]
        metadata["page_count"] = len(doc)
        metadata["title"] = doc.metadata.get("title", "")
        doc.close()
        return "\n\n".join(pages), metadata
    finally:
        os.remove(tmp_path)


def _extract_docx(raw_bytes: bytes, metadata: dict) -> tuple[str, dict]:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(raw_bytes)
        tmp_path = tmp.name
    try:
        doc = DocxDocument(tmp_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        metadata["paragraph_count"] = len(paragraphs)
        return "\n\n".join(paragraphs), metadata
    finally:
        os.remove(tmp_path)


def _extract_pptx(raw_bytes: bytes, metadata: dict) -> tuple[str, dict]:
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        tmp.write(raw_bytes)
        tmp_path = tmp.name
    try:
        prs = Presentation(tmp_path)
        slides_text = [
            " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
            for slide in prs.slides
        ]
        metadata["slide_count"] = len(prs.slides)
        return "\n\n".join(slides_text), metadata
    finally:
        os.remove(tmp_path)
